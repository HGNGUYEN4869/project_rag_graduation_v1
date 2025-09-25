import os
import re
import uuid
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document
from typing import Dict, Any
from enviroment.envGlobal import DB_JSON
from transcriptToJson import write_next_json

# =========================
# Utility
# =========================
def clean_text(text: str) -> str:
    text = re.sub(r"[^0-9A-Za-zÀ-ỹ\s\.,:;()\-\+]", "", text)
    return re.sub(r"\s+", " ", text).strip()

def normalize_book_name(file_path: str) -> str:
    name = os.path.splitext(os.path.basename(file_path))[0]
    name = re.sub(r"[^0-9A-Za-zÀ-ỹ\s]", "", name)
    name = re.sub(r"\s+", "_", name)
    return name.strip("_")

def save_json(data: Dict[str, Any], path: str):
    import json
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[+] Saved to {path}")

# =========================
# PDF Helpers
# =========================
def extract_page_elements(page, path, page_num):
    elements = []
    text = page.extract_text(x_tolerance=2, y_tolerance=2)

    if text:
        for line in text.split("\n"):
            if re.match(r"^\s*trang\s*\d+\s*$", line.strip(), re.IGNORECASE):
                continue
            bbox = page.chars[0] if page.chars else {"top": 0}
            elements.append({
                "id": str(uuid.uuid4()),
                "text": line.strip(),
                "top": bbox["top"] if "top" in bbox else 0,
                "metadata": {"page": page_num, "source": path, "type": "text"}
            })
    else:
        pil_image = page.to_image(resolution=300).original
        ocr_text = pytesseract.image_to_string(pil_image, lang="vie+eng")
        for i, line in enumerate(ocr_text.split("\n")):
            if re.match(r"^\s*trang\s*\d+\s*$", line.strip(), re.IGNORECASE):
                continue
            elements.append({
                "id": str(uuid.uuid4()),
                "text": line.strip(),
                "top": i * 20,
                "metadata": {"page": page_num, "source": path, "type": "ocr_text"}
            })

    elements.sort(key=lambda x: x["top"])
    return elements

def extract_images_from_pdf(page, page_num, out_dir="images"):
    os.makedirs(out_dir, exist_ok=True)
    pil_image = page.to_image(resolution=300).original
    img_path = os.path.join(out_dir, f"pdf_page_{page_num}.png")
    pil_image.save(img_path)
    return img_path

# =========================
# DOCX Helpers
# =========================
def extract_images_from_docx(doc_path, out_dir="images"):
    os.makedirs(out_dir, exist_ok=True)
    doc = Document(doc_path)
    image_paths = []

    rels = doc.part._rels
    for rel in rels:
        rel = rels[rel]
        if "image" in rel.target_ref:
            img_data = rel.target_part.blob
            img_name = f"docx_img_{len(image_paths)}.png"
            img_path = os.path.join(out_dir, img_name)
            with open(img_path, "wb") as f:
                f.write(img_data)
            image_paths.append(img_path)
    return image_paths

# =========================
# Main loader
# =========================
def load_document(file_path: str) -> Dict[str, Any]:
    book_name = normalize_book_name(file_path)
    data = {"text": [], "images": []}

    if file_path.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                elements = extract_page_elements(page, file_path, page_num)

                merged_content = []
                buffer = ""
                prev_top = None
                for el in elements:
                    text = clean_text(el["text"])
                    if not text:
                        continue
                    top = el["top"]
                    if prev_top is not None and (top - prev_top > 15 or buffer.endswith(('.', ':', ';'))):
                        merged_content.append(buffer.strip())
                        buffer = text
                    else:
                        buffer += " " + text
                    prev_top = top
                if buffer:
                    merged_content.append(buffer.strip())

                if merged_content:
                    data["text"].append({"page": page_num, "content": merged_content})

                img_path = extract_images_from_pdf(page, page_num)
                ocr_text = pytesseract.image_to_string(Image.open(img_path), lang="vie+eng")
                data["images"].append({
                    "path": img_path,
                    "ocr_text": ocr_text,
                    "chart_data": {},
                    "table_data": []
                })

    elif file_path.endswith(".docx"):
        doc = Document(file_path)
        merged_content = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        if merged_content:
            data["text"].append({"page": 1, "content": merged_content})

        image_paths = extract_images_from_docx(file_path)
        for img in image_paths:
            ocr_text = pytesseract.image_to_string(Image.open(img), lang="vie+eng")
            data["images"].append({
                "path": img,
                "ocr_text": ocr_text,
                "chart_data": {},
                "table_data": []
            })

    else:
        raise ValueError("Unsupported file format: must be .pdf or .docx")

    final_data = {book_name: data}
    write_next_json(final_data, DB_JSON)
    os.startfile(DB_JSON)

